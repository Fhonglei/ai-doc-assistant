"""Document parser: PDF (PyPDF2/pdfplumber), DOCX (python-docx), TXT (chardet)."""

from dataclasses import dataclass, field
from typing import List, Optional
from app.utils.file_utils import detect_encoding


@dataclass
class PageInfo:
    page_number: int
    text: str


@dataclass
class ParsedDocument:
    text: str
    pages: List[PageInfo] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)


def parse_document(content: bytes, filename: str, file_type: str) -> ParsedDocument:
    """Dispatch to the appropriate parser based on file type."""
    if file_type == "pdf":
        return _parse_pdf(content, filename)
    elif file_type == "docx":
        return _parse_docx(content, filename)
    elif file_type == "txt":
        return _parse_txt(content, filename)
    else:
        raise ValueError(f"UNSUPPORTED_FILE_TYPE: Unknown file type '{file_type}'.")


def _parse_pdf(content: bytes, filename: str) -> ParsedDocument:
    """Parse PDF with PyPDF2, falling back to pdfplumber."""
    import io
    from PyPDF2 import PdfReader

    pages: List[PageInfo] = []
    full_text_parts: List[str] = []

    try:
        reader = PdfReader(io.BytesIO(content))
        for i, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            pages.append(PageInfo(page_number=i + 1, text=text))
            full_text_parts.append(text)
    except Exception as e:
        # Try pdfplumber fallback
        try:
            return _parse_pdf_fallback(content, filename)
        except ImportError:
            raise ValueError(
                f"CORRUPTED_FILE: Could not read PDF file. "
                f"PyPDF2 error: {str(e)[:200]}"
            )

    full_text = "\n\n".join(full_text_parts)

    # Detect scanned PDF (average < 20 chars per page)
    if pages and sum(len(p.text) for p in pages) / len(pages) < 20:
        raise ValueError(
            "SCANNED_PDF: This PDF appears to be scanned (image-only). "
            "OCR functionality is not yet available."
        )

    if not full_text.strip():
        raise ValueError("EMPTY_DOCUMENT: No extractable text found in this document.")

    return ParsedDocument(
        text=full_text,
        pages=pages,
        metadata={"filename": filename, "file_type": "pdf", "page_count": len(pages)},
    )


def _parse_pdf_fallback(content: bytes, filename: str) -> ParsedDocument:
    """Fallback PDF parser using pdfplumber."""
    import io
    import pdfplumber

    pages: List[PageInfo] = []
    full_text_parts: List[str] = []

    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            pages.append(PageInfo(page_number=i + 1, text=text))
            full_text_parts.append(text)

    full_text = "\n\n".join(full_text_parts)

    if not full_text.strip():
        raise ValueError("EMPTY_DOCUMENT: No extractable text found in this document.")

    return ParsedDocument(
        text=full_text,
        pages=pages,
        metadata={"filename": filename, "file_type": "pdf", "page_count": len(pages)},
    )


def _parse_docx(content: bytes, filename: str) -> ParsedDocument:
    """Parse DOCX using python-docx."""
    import io
    from docx import Document

    try:
        doc = Document(io.BytesIO(content))
    except Exception as e:
        raise ValueError(
            f"CORRUPTED_FILE: Could not read DOCX file. It may be corrupted. "
            f"Error: {str(e)[:200]}"
        )

    paragraphs = []
    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text)

    # Also extract table text
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text for cell in row.cells if cell.text.strip())
            if row_text.strip():
                paragraphs.append(row_text)

    full_text = "\n\n".join(paragraphs)

    if not full_text.strip():
        raise ValueError("EMPTY_DOCUMENT: No extractable text found in this document.")

    # DOCX doesn't have native page numbers
    return ParsedDocument(
        text=full_text,
        pages=[PageInfo(page_number=1, text=full_text)],
        metadata={"filename": filename, "file_type": "docx", "paragraph_count": len(paragraphs)},
    )


def _parse_txt(content: bytes, filename: str) -> ParsedDocument:
    """Parse TXT with encoding detection."""
    encoding = detect_encoding(content)
    try:
        text = content.decode(encoding)
    except UnicodeDecodeError:
        text = content.decode("latin-1")

    if not text.strip():
        raise ValueError("EMPTY_DOCUMENT: No extractable text found in this document.")

    return ParsedDocument(
        text=text,
        pages=[PageInfo(page_number=1, text=text)],
        metadata={"filename": filename, "file_type": "txt", "encoding": encoding},
    )
